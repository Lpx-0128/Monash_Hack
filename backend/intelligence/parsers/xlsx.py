"""XLSX parsing over real typed cells.

Conventions:

* ``sheet_cell`` carries the exact sheet name and A1 coordinate.
* A merged range is read once, at its top-left origin.
* Missing, zero, False and empty-string stay distinct: ``cell.value or ""``
  is never used.
* ``data_only`` exposes the value Excel last stored. openpyxl does not calculate
  formulas, so a formula with no cached value is reported, not treated as blank.
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

from ..config import PARSER_VERSIONS
from ..types import (
    Block,
    Diagnostic,
    ParsedDocument,
    ParseStatus,
    PermanentProcessingError,
    SheetCell,
)
from .base import iter_label_spans, normalize_line_endings

NAME = "xlsx"
VERSION = PARSER_VERSIONS["xlsx"]

FORMULA_WITHOUT_CACHED_VALUE = "FORMULA_WITHOUT_CACHED_VALUE"


def render_cell(value) -> str:
    """The parser's persisted rendering of a typed cell.

    ``sheet_cell`` evidence quotes must equal this rendering. The underlying
    scalar and type stay available in the block context.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        exact = Decimal(str(value))
        normalized = exact.normalize()
        if normalized == normalized.to_integral_value():
            return str(int(normalized))
        return format(normalized, "f")
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return normalize_line_endings(str(value))


def _cell_type(value) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return "datetime"
    return "text"


def parse(data: bytes, *, document_id: str, content_hash: str,
          filename: str, media_type: str, max_cells: int = 200_000) -> ParsedDocument:
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise PermanentProcessingError(f"the XLSX parser dependency is unavailable: {exc}") from exc

    def load(data_only: bool):
        return openpyxl.load_workbook(io.BytesIO(data), data_only=data_only, read_only=False)

    try:
        values_wb = load(True)
        formulas_wb = load(False)
    except (InvalidFileException, KeyError, ValueError) as exc:
        return ParsedDocument(
            document_id=document_id, content_hash=content_hash, parser_name=NAME,
            parser_version=VERSION, status=ParseStatus.UNREADABLE,
            diagnostics=(Diagnostic("CORRUPT_SOURCE", "the workbook could not be opened",
                                    {"error": str(exc)}),),
            media_type=media_type, filename=filename,
        )
    except Exception as exc:
        raise PermanentProcessingError(f"unexpected XLSX parser failure: {exc!r}") from exc

    diagnostics: list[Diagnostic] = []
    blocks: list[Block] = []
    text_parts: list[str] = []
    order = 0
    cell_budget = max_cells

    for sheet_name in values_wb.sheetnames:
        values_ws = values_wb[sheet_name]
        formulas_ws = formulas_wb[sheet_name]

        # Merged ranges are read once, at their top-left origin.
        merged_origin: dict[tuple[int, int], tuple[int, int]] = {}
        for cell_range in values_ws.merged_cells.ranges:
            origin = (cell_range.min_row, cell_range.min_col)
            for row in range(cell_range.min_row, cell_range.max_row + 1):
                for column in range(cell_range.min_col, cell_range.max_col + 1):
                    merged_origin[(row, column)] = origin

        rendered: dict[tuple[int, int], str] = {}
        typed: dict[tuple[int, int], tuple[str, object]] = {}
        for row in values_ws.iter_rows():
            for cell in row:
                position = (cell.row, cell.column)
                if merged_origin.get(position, position) != position:
                    continue  # merged continuation: never counted twice
                cell_budget -= 1
                if cell_budget < 0:
                    diagnostics.append(
                        Diagnostic("CELL_LIMIT", "the configured cell budget was reached",
                                   {"sheet": sheet_name, "limit": max_cells})
                    )
                    break
                value = cell.value
                formula = formulas_ws.cell(row=cell.row, column=cell.column).value
                if value is None and isinstance(formula, str) and formula.startswith("="):
                    diagnostics.append(
                        Diagnostic(
                            FORMULA_WITHOUT_CACHED_VALUE,
                            "the workbook stores a formula with no cached value; "
                            "openpyxl does not calculate it",
                            {"sheet": sheet_name, "cell": cell.coordinate, "formula": formula},
                        )
                    )
                rendered[position] = render_cell(value)
                typed[position] = (_cell_type(value), value)
            if cell_budget < 0:
                break

        for (row, column), content in sorted(rendered.items()):
            if not content.strip():
                continue
            text_parts.append(content)
            coordinate = f"{get_column_letter(column)}{row}"
            kind, raw_value = typed[(row, column)]

            spans = list(iter_label_spans(content))
            if spans:
                for span in spans:
                    blocks.append(Block(
                        block_id=f"{sheet_name}!{coordinate}#{order:04d}",
                        label=span.label,
                        value_text=span.value_text,
                        value_locator=SheetCell(sheet_name, coordinate),
                        label_locator=SheetCell(sheet_name, coordinate),
                        context={"sheet": sheet_name, "cell": coordinate,
                                 "cell_type": kind, "in_cell": True},
                        order=order,
                    ))
                    order += 1
                continue

            label = content.strip().rstrip(":").strip()
            if not label:
                continue
            # Value cells are the non-empty cells to the right of the label, in the
            # same row. Pairing downwards is deliberately not done: it invents a
            # label/value relationship the sheet structure does not express.
            followers = [
                (r, c) for (r, c) in sorted(rendered)
                if r == row and c > column and rendered[(r, c)].strip()
            ]
            if not followers:
                continue
            first = followers[0]
            first_coordinate = f"{get_column_letter(first[1])}{first[0]}"
            extra = followers[1:2]  # an adjacent continuation column, e.g. an address
            blocks.append(Block(
                block_id=f"{sheet_name}!{coordinate}#{order:04d}",
                label=label,
                value_text=rendered[first],
                value_locator=SheetCell(sheet_name, first_coordinate),
                label_locator=SheetCell(sheet_name, coordinate),
                extra_value_locators=tuple(
                    SheetCell(sheet_name, f"{get_column_letter(c)}{r}") for r, c in extra
                ),
                context={
                    "sheet": sheet_name,
                    "cell": first_coordinate,
                    "cell_type": typed[first][0],
                    "cell_value": typed[first][1] if typed[first][0] in ("int", "float") else None,
                    "extra_value_texts": [rendered[position] for position in extra],
                },
                order=order,
            ))
            order += 1

    if not blocks and not any(part.strip() for part in text_parts):
        diagnostics.append(Diagnostic("EMPTY_DOCUMENT", "the workbook carried no cell content", {}))
        return ParsedDocument(
            document_id=document_id, content_hash=content_hash, parser_name=NAME,
            parser_version=VERSION, status=ParseStatus.UNREADABLE,
            diagnostics=tuple(diagnostics), media_type=media_type, filename=filename,
        )

    return ParsedDocument(
        document_id=document_id,
        content_hash=content_hash,
        parser_name=NAME,
        parser_version=VERSION,
        status=ParseStatus.OK,
        text="\n".join(text_parts),
        blocks=tuple(blocks),
        diagnostics=tuple(diagnostics),
        media_type=media_type,
        filename=filename,
    )
