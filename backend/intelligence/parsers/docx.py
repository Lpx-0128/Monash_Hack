"""DOCX parsing over body paragraphs and tables in document order.

Indexing conventions relied upon by G1:

* ``docx_paragraph.index`` is zero-based over body paragraphs in body order.
* ``docx_table_cell`` uses zero-based table/row/column. A merged cell is counted
  once, at its canonical origin — the first (row, column) at which the
  underlying cell appears in row-major order.
"""

from __future__ import annotations

import io

from ..config import PARSER_VERSIONS
from ..types import (
    Block,
    Diagnostic,
    DocxParagraph,
    DocxTableCell,
    ParsedDocument,
    ParseStatus,
    PermanentProcessingError,
)
from .base import iter_label_spans, normalize_line_endings

NAME = "docx"
VERSION = PARSER_VERSIONS["docx"]


def parse(data: bytes, *, document_id: str, content_hash: str,
          filename: str, media_type: str) -> ParsedDocument:
    try:
        import docx as python_docx
        from docx.opc.exceptions import PackageNotFoundError
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise PermanentProcessingError(f"the DOCX parser dependency is unavailable: {exc}") from exc

    try:
        document = python_docx.Document(io.BytesIO(data))
    except PackageNotFoundError as exc:
        return ParsedDocument(
            document_id=document_id, content_hash=content_hash, parser_name=NAME,
            parser_version=VERSION, status=ParseStatus.UNREADABLE,
            diagnostics=(Diagnostic("CORRUPT_SOURCE", "the DOCX package could not be opened",
                                    {"error": str(exc)}),),
            media_type=media_type, filename=filename,
        )
    except Exception as exc:
        raise PermanentProcessingError(f"unexpected DOCX parser failure: {exc!r}") from exc

    body = document.element.body
    diagnostics: list[Diagnostic] = []
    blocks: list[Block] = []
    text_parts: list[str] = []
    paragraph_index = 0
    table_index = 0
    order = 0

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            content = normalize_line_endings(paragraph.text)
            text_parts.append(content)
            for span in iter_label_spans(content):
                blocks.append(Block(
                    block_id=f"para{paragraph_index:04d}_{order:04d}",
                    label=span.label,
                    value_text=span.value_text,
                    value_locator=DocxParagraph(paragraph_index),
                    label_locator=DocxParagraph(paragraph_index),
                    context={"paragraph": paragraph_index},
                    order=order,
                ))
                order += 1
            paragraph_index += 1
        elif tag == "tbl":
            table = Table(child, document)
            # A merged cell is reached from several grid positions. Key it by the
            # underlying element's document path — a stable identity, unlike id(),
            # which lxml recycles when a proxy is collected.
            seen_cells: dict[str, tuple[int, int]] = {}
            grid: list[list[tuple[int, int, str]]] = []
            for row_index, row in enumerate(table.rows):
                row_cells: list[tuple[int, int, str]] = []
                for column_index, cell in enumerate(row.cells):
                    element = cell._tc
                    key = element.getroottree().getpath(element)
                    origin = seen_cells.setdefault(key, (row_index, column_index))
                    if origin != (row_index, column_index):
                        continue  # merged continuation: never counted twice
                    row_cells.append((row_index, column_index, normalize_line_endings(cell.text)))
                grid.append(row_cells)
                text_parts.extend(value for _, _, value in row_cells)

            for row_cells in grid:
                for position, (row_index, column_index, content) in enumerate(row_cells):
                    # A label cell paired with the adjacent contributing value cells.
                    spans = list(iter_label_spans(content))
                    if spans:
                        for span in spans:
                            blocks.append(Block(
                                block_id=f"tbl{table_index}_{row_index}_{column_index}_{order:04d}",
                                label=span.label,
                                value_text=span.value_text,
                                value_locator=DocxTableCell(table_index, row_index, column_index),
                                label_locator=DocxTableCell(table_index, row_index, column_index),
                                context={"table": table_index, "row": row_index,
                                         "column": column_index, "in_cell": True},
                                order=order,
                            ))
                            order += 1
                        continue

                    label = content.strip().rstrip(":").strip()
                    if not label:
                        continue
                    followers = row_cells[position + 1:]
                    values = [(r, c, v) for r, c, v in followers if v.strip()]
                    if not values:
                        continue
                    # Keep every contributing value cell as its own evidence reference
                    # rather than manufacturing one concatenated quote.
                    first_row, first_col, first_value = values[0]
                    blocks.append(Block(
                        block_id=f"tbl{table_index}_{row_index}_{column_index}_{order:04d}",
                        label=label,
                        value_text=first_value,
                        value_locator=DocxTableCell(table_index, first_row, first_col),
                        label_locator=DocxTableCell(table_index, row_index, column_index),
                        extra_value_locators=tuple(
                            DocxTableCell(table_index, r, c) for r, c, _ in values[1:]
                        ),
                        context={
                            "table": table_index,
                            "row": row_index,
                            "label_column": column_index,
                            "extra_value_texts": [v for _, _, v in values[1:]],
                        },
                        order=order,
                    ))
                    order += 1
            table_index += 1

    if not blocks and not any(part.strip() for part in text_parts):
        diagnostics.append(Diagnostic("EMPTY_DOCUMENT", "the DOCX body carried no text", {}))
        return ParsedDocument(
            document_id=document_id, content_hash=content_hash, parser_name=NAME,
            parser_version=VERSION, status=ParseStatus.UNREADABLE,
            diagnostics=tuple(diagnostics), media_type=media_type, filename=filename,
        )

    return ParsedDocument(
        document_id=document_id,
        content_hash=content_hash,
        parser_name=NAME,
        parser_version=VERSION,
        status=ParseStatus.OK,
        text="\n".join(text_parts),
        blocks=tuple(blocks),
        diagnostics=tuple(diagnostics),
        media_type=media_type,
        filename=filename,
    )
