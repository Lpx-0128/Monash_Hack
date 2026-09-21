"""Parser behaviour and stable evidence (handoff cases A-11 to A-18)."""

import pytest

from backend.intelligence.parsers import docx as docx_parser
from backend.intelligence.parsers import pdf as pdf_parser
from backend.intelligence.parsers.base import iter_label_spans, normalize_line_endings
from backend.intelligence.parsers.router import detect_format
from backend.intelligence.parsers.xlsx import render_cell
from backend.intelligence.policies import aliases
from backend.intelligence.types import (
    DocxTableCell,
    ParseStatus,
    PermanentProcessingError,
    SheetCell,
    TextRange,
)

from .conftest import DEMO_ATTACHMENTS, PARTICIPANT_ROOT

TXT_WITH_ADDRESS = (
    "SHIPPING INSTRUCTION\n"
    "========================================\n"
    "\n"
    "Shipper: APRIL FAR EAST (M) SDN BHD\n"
    "  TOWER 2, AVENUE 5, LEVEL 6; BANGSAR SOUTH CITY\n"
    "  TEL: +60 3 1234 5678\n"
    "  P.O. BOX: 8899, KUALA LUMPUR\n"
    "Consignee: EAST BRIGHT FZ-LLC\n"
    "NET WEIGHT: _______ MTS\n"
    "Gross Weight (KG): 18500 KG\n"
)


def test_a11_multiline_party_block_is_kept_whole_with_exact_offsets(parse):
    """A-11: TEL and P.O. BOX lines stay inside the party block, offsets exact."""
    document = parse(TXT_WITH_ADDRESS.encode("utf-8"))
    shipper = next(b for b in document.blocks if b.label == "Shipper")

    assert "TOWER 2, AVENUE 5" in shipper.value_text
    assert "TEL: +60 3 1234 5678" in shipper.value_text
    assert "P.O. BOX: 8899, KUALA LUMPUR" in shipper.value_text
    # The block stops at the next boundary label.
    assert "EAST BRIGHT" not in shipper.value_text

    locator = shipper.value_locator
    assert isinstance(locator, TextRange)
    assert document.text[locator.start:locator.end] == shipper.value_text
    assert document.locator_texts(locator) == (shipper.value_text,)


def test_a12_empty_value_never_borrows_the_next_label(parse):
    """A-12: an empty labelled field stays empty."""
    source = "Gross Weight (KG):\nPort of Loading: SINGAPORE\n"
    document = parse(source.encode("utf-8"))
    weight = next(b for b in document.blocks if b.label == "Gross Weight (KG)")
    assert weight.value_text.strip() == ""
    assert "SINGAPORE" not in weight.value_text


def test_placeholder_value_is_kept_verbatim(parse):
    """The raw placeholder token survives, so the question can quote it."""
    document = parse(TXT_WITH_ADDRESS.encode("utf-8"))
    net = next(b for b in document.blocks if b.label == "NET WEIGHT")
    assert net.value_text.strip() == "_______ MTS"
    # A net weight is never a source for gross weight.
    assert aliases.match_field("NET WEIGHT") is None


def test_line_endings_are_normalized_and_offsets_follow_the_result(parse):
    document = parse("Shipper: ACME\r\nConsignee: BETA\r\n".encode("utf-8"))
    assert "\r" not in document.text
    for block in document.blocks:
        start, end = block.value_locator.start, block.value_locator.end
        assert document.text[start:end] == block.value_text


def test_undecodable_text_is_a_source_issue_not_a_corrupted_field(parse):
    """errors='replace' is never used to fabricate a readable value."""
    document = parse(b"Shipper: \xff\xfe\x00\x00\x80ACME", filename="broken.txt")
    assert document.status in (ParseStatus.UNREADABLE, ParseStatus.UNSUPPORTED)


def test_a13_docx_bilingual_labels_and_multiline_cells(parse_demo):
    """A-13: bilingual table labels resolve and address cells stay complete."""
    document = parse_demo("demo_office_BL.docx")
    assert document.status is ParseStatus.OK

    shipper = next(b for b in document.blocks if aliases.match_field(b.label) == "shipper")
    assert isinstance(shipper.value_locator, DocxTableCell)
    assert "NORTHWIND PAPER EXPORTS PTE LTD" in shipper.value_text
    # The address after the line break is not discarded.
    assert "12 KEPPEL ROAD" in shipper.value_text

    resolved = {aliases.match_field(b.label) for b in document.blocks}
    for field in ("shipper", "consignee", "notify_party", "port_of_loading",
                  "port_of_discharge", "container_count", "gross_weight_kg"):
        assert field in resolved, field


def test_a14_merged_docx_cells_are_counted_once(parse):
    """A-14: a merged cell appears at one canonical origin, never twice."""
    import io

    import docx

    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "MERGED HEADING"
    table.cell(1, 0).text = "Container Count"
    table.cell(1, 1).text = "4 x 40'HC"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse(buffer.getvalue(), filename="merged.docx")
    merged_quotes = [b for b in parsed.blocks if b.value_text == "MERGED HEADING"]
    assert len(merged_quotes) <= 1

    count = next(b for b in parsed.blocks if aliases.match_field(b.label) == "container_count")
    assert count.value_text == "4 x 40'HC"


def test_a15_xlsx_keeps_zero_blank_false_and_text_distinct():
    """A-15: `cell.value or ""` is never used, so the distinctions survive."""
    assert render_cell(0) == "0"
    assert render_cell(None) == ""
    assert render_cell(False) == "FALSE"
    assert render_cell(True) == "TRUE"
    assert render_cell("") == ""
    assert render_cell(0.0) == "0"
    assert render_cell(12.5) == "12.5"
    # Zero is a present value, not a missing one.
    assert render_cell(0) != render_cell(None)


def test_a15_formula_without_a_cached_value_is_diagnosed(parse):
    """openpyxl does not calculate: a missing cached value is reported, not blank."""
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S.I."
    sheet["A1"] = "Gross Weight (KG)"
    sheet["B1"] = "=SUM(C1:C5)"
    buffer = io.BytesIO()
    workbook.save(buffer)

    document = parse(buffer.getvalue(), filename="formula.xlsx")
    codes = {d.code for d in document.diagnostics}
    assert "FORMULA_WITHOUT_CACHED_VALUE" in codes


def test_xlsx_cells_use_exact_sheet_names_and_a1_coordinates(parse_demo):
    document = parse_demo("demo_office_SI.xlsx")
    weight = next(b for b in document.blocks if aliases.match_field(b.label) == "gross_weight_kg")
    assert isinstance(weight.value_locator, SheetCell)
    assert weight.value_locator.sheet == "S.I."
    assert weight.value_locator.cell[0].isalpha()
    # The evidence quote equals the parser's persisted cell rendering.
    assert document.locator_texts(weight.value_locator)[0] == weight.value_text


def test_a16_pdf_total_is_available_alongside_its_detail_rows(parse_demo):
    """A-16: a declared total is a candidate; detail rows are not summed into it."""
    document = parse_demo("demo_pdf_BL.pdf")
    assert document.status is ParseStatus.OK

    totals = [b for b in document.blocks if aliases.match_field(b.label) == "gross_weight_kg"]
    assert len(totals) == 1
    assert totals[0].value_text.strip() == "21600 KG"
    assert aliases.is_total_label(totals[0].label)
    # Each row's 7200 is in the page text but carries no gross-weight label.
    assert "7200" in document.pages[0]


def test_a17_image_only_pdf_is_not_a_successful_extraction(parse_demo):
    """A-17: no text layer is a readability limit, not a clean blank form."""
    document = parse_demo("demo_scan_BL.pdf")
    assert document.status is ParseStatus.UNREADABLE
    codes = {d.code for d in document.diagnostics}
    assert pdf_parser.IMAGE_ONLY_PAGE in codes
    assert pdf_parser.IMAGE_ONLY_DOCUMENT in codes
    assert document.blocks == ()


def test_a18_corrupt_pdf_is_a_source_issue(parse_demo):
    """A-18: an unopenable PDF is established source corruption."""
    document = parse_demo("demo_corrupt_BL.pdf")
    assert document.status is ParseStatus.UNREADABLE
    assert {d.code for d in document.diagnostics} == {pdf_parser.CORRUPT_SOURCE}


def test_a18_an_unexpected_parser_exception_stays_a_technical_failure(monkeypatch, parse_demo):
    """A-18: an unknown error is never dressed up as an unreadable document."""
    import pypdf

    class Exploding:
        def __init__(self, *_args, **_kwargs):
            raise MemoryError("simulated internal failure")

    monkeypatch.setattr(pypdf, "PdfReader", Exploding)
    with pytest.raises(PermanentProcessingError):
        parse_demo("demo_pdf_SI.pdf")


@pytest.mark.parametrize("name,expected", [
    ("demo_match_SI.txt", "text"),
    ("demo_pdf_SI.pdf", "pdf"),
    ("demo_office_BL.docx", "docx"),
    ("demo_office_SI.xlsx", "xlsx"),
])
def test_router_dispatches_on_content_signature(name, expected):
    data = (DEMO_ATTACHMENTS / name).read_bytes()
    # The extension is deliberately wrong; the signature decides.
    assert detect_format(data, "misleading.txt") == expected


def test_router_refuses_an_arbitrary_zip():
    """A ZIP is not automatically a Word or Excel document."""
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hello")
    assert detect_format(buffer.getvalue(), "thing.docx") == "unsupported"


def test_unsupported_content_is_reported_not_guessed(parse):
    document = parse(b"\x00\x01\x02\x03binary-nonsense\x00", filename="thing.bin")
    assert document.status is ParseStatus.UNSUPPORTED


def test_parse_is_reproducible_for_the_same_bytes(parse_demo):
    """The same bytes and parser version yield identical evidence."""
    first = parse_demo("demo_match_SI.txt")
    second = parse_demo("demo_match_SI.txt")
    assert first.content_hash == second.content_hash
    assert [b.value_text for b in first.blocks] == [b.value_text for b in second.blocks]
    assert [b.value_locator for b in first.blocks] == [b.value_locator for b in second.blocks]


def test_real_participant_formats_all_parse(parse):
    """Every required format is exercised against the actual bundle."""
    samples = {
        "email_004_SI.txt": ParseStatus.OK,
        "email_059_BL.pdf": ParseStatus.OK,
        "email_107_BL.docx": ParseStatus.OK,
        "email_005_SI.xlsx": ParseStatus.OK,
        "email_512_BL.pdf": ParseStatus.UNREADABLE,
        "email_511_BL.pdf": ParseStatus.UNREADABLE,
    }
    for name, expected in samples.items():
        path = PARTICIPANT_ROOT / "attachments" / name
        document = parse(path.read_bytes(), filename=name)
        assert document.status is expected, f"{name} parsed as {document.status}"
