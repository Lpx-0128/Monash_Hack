#!/usr/bin/env python3
"""Build the synthetic demo source registry.

These documents are written for this project. They contain no participant data,
so they are the only sources a public DEMO caller may stream. Expected values
live in the tests and were read off these documents by hand, never produced by
the extractor under test.

    python scripts/build_demo_fixtures.py [--root resources/demo-fixtures]
"""

from __future__ import annotations

import argparse
import json
import shutil
import zlib
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_ROOT = REPO_ROOT / "resources" / "demo-fixtures"


# ---------------------------------------------------------------------------
# Minimal PDF writer: a real text layer, no third-party dependency.
# ---------------------------------------------------------------------------

def _escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def write_text_pdf(path: Path, lines: list[str]) -> None:
    """A single-page PDF whose text layer contains ``lines``."""
    content = ["BT", "/F1 11 Tf", "14 TL", "40 760 Td"]
    for line in lines:
        content.append(f"({_escape(line)}) Tj")
        content.append("T*")
    content.append("ET")
    stream = "\n".join(content).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    _write_pdf(path, objects)


def write_image_only_pdf(path: Path) -> None:
    """A page carrying an image XObject and no text layer at all."""
    # A 2x2 greyscale image, Flate-compressed.
    raw = bytes([0, 255, 255, 0])
    image = zlib.compress(raw)
    stream = b"q 200 0 0 120 40 640 cm /Im1 Do Q"

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /XObject << /Im1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceGray "
        b"/BitsPerComponent 8 /Filter /FlateDecode /Length " + str(len(image)).encode()
        + b" >>\nstream\n" + image + b"\nendstream",
    ]
    _write_pdf(path, objects)


def _write_pdf(path: Path, objects: list[bytes]) -> None:
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_at}\n%%EOF\n").encode()
    path.write_bytes(bytes(out))


def write_corrupt_pdf(path: Path) -> None:
    """Claims to be a PDF and then stops. An established source corruption."""
    path.write_bytes(b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")


# ---------------------------------------------------------------------------
# DOCX and XLSX
# ---------------------------------------------------------------------------

def write_docx(path: Path, heading: str, rows: list[tuple[str, str]]) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph(heading)
    table = document.add_table(rows=len(rows), cols=2)
    for index, (label, value) in enumerate(rows):
        table.cell(index, 0).text = label
        cell = table.cell(index, 1)
        parts = value.split("\n")
        cell.text = parts[0]
        for extra in parts[1:]:
            cell.add_paragraph(extra)
    document.save(path)


def write_xlsx(path: Path, heading: str, rows: list[tuple[str, object]],
               sheet_name: str = "S.I.") -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet["A1"] = heading
    for index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=index, column=1, value=label)
        sheet.cell(row=index, column=2, value=value)
    workbook.save(path)


# ---------------------------------------------------------------------------
# Document bodies
# ---------------------------------------------------------------------------

SI_HEADING = "SHIPPING INSTRUCTION\n========================================\n"
BL_HEADING = "BILL OF LADING (DRAFT)\n========================================\n"

BASE_SI = """Shipper: NORTHWIND PAPER EXPORTS PTE LTD
  12 KEPPEL ROAD, #08-02; SINGAPORE 089057
Consignee: MERIDIAN TRADING GMBH
  HAFENSTRASSE 44; 20457 HAMBURG, GERMANY
Notify Party: MERIDIAN TRADING GMBH
Port of Loading (POL): SINGAPORE (SGSIN)
Port of Discharge (POD): HAMBURG, GERMANY (DEHAM)
Total Containers: 4 x 40'HC
Gross Weight (KG): 18500 KG
Vessel: CORAL STAR V.118E
Booking Ref: DEMO-BK-00041
Freight: PREPAID
"""

BASE_BL = """SHIPPER: NORTHWIND PAPER EXPORTS PTE LTD
  12 KEPPEL ROAD, #08-02; SINGAPORE 089057
Consignee (Non-Negotiable): MERIDIAN TRADING GMBH
  HAFENSTRASSE 44; 20457 HAMBURG, GERMANY
NOTIFY PARTY: MERIDIAN TRADING GMBH
Load Port: SINGAPORE (SGSIN)
Discharge Port: HAMBURG, GERMANY (DEHAM)
Container Count: 4 x 40'HC
Gross Wt (kgs): 18500 KG
Export Carrier (vessel, voyage): CORAL STAR V.118E
Bill of Lading No.: DEMO-BL-77120
Freight: PREPAID
"""

PACKING_LIST = """PACKING LIST
========================================

Packing List No.: DEMO-PL-9001
Seller: NORTHWIND PAPER EXPORTS PTE LTD
Buyer: MERIDIAN TRADING GMBH
Total Packages: 820 CARTONS
NET WEIGHT: 17900 KG
Marks and Numbers: DEMO/HAM/2026

*** THIS IS A PACKING LIST - NOT A BILL OF LADING ***
"""


def build(root: Path) -> None:
    if root.exists():
        shutil.rmtree(root)
    inbox = root / "inbox"
    attachments = root / "attachments"
    inbox.mkdir(parents=True)
    attachments.mkdir(parents=True)

    emails: list[dict] = []

    def email(email_id: str, subject: str, body: str, files: list[str]) -> None:
        emails.append({
            "email_id": email_id,
            "from": "demo.sender@example.invalid",
            "subject": subject,
            "body": body,
            "attachments": [f"attachments/{name}" for name in files],
        })

    compare_body = ("Dear Team,\n\nPlease compare the SI and draft BL and confirm "
                    "the shipment details agree.\n\nRegards,\nDemo Desk\n")

    # 1. Every field agrees.
    (attachments / "demo_match_SI.txt").write_text(SI_HEADING + BASE_SI, encoding="utf-8")
    (attachments / "demo_match_BL.txt").write_text(BL_HEADING + BASE_BL, encoding="utf-8")
    email("email_demo_match", "Please check SI vs draft BL - DEMO-BK-00041",
          compare_body, ["demo_match_SI.txt", "demo_match_BL.txt"])

    # 2. One confirmed difference: the consignee.
    (attachments / "demo_mismatch_SI.txt").write_text(SI_HEADING + BASE_SI, encoding="utf-8")
    (attachments / "demo_mismatch_BL.txt").write_text(
        BL_HEADING + BASE_BL.replace("Consignee (Non-Negotiable): MERIDIAN TRADING GMBH",
                                     "Consignee (Non-Negotiable): ORION IMPORTS SARL"),
        encoding="utf-8")
    email("email_demo_mismatch", "Please check SI vs draft BL - DEMO-BK-00042",
          compare_body, ["demo_mismatch_SI.txt", "demo_mismatch_BL.txt"])

    # 3. The BL gross weight is blank: a targeted value question.
    (attachments / "demo_missing_SI.txt").write_text(SI_HEADING + BASE_SI, encoding="utf-8")
    (attachments / "demo_missing_BL.txt").write_text(
        BL_HEADING + BASE_BL.replace("Gross Wt (kgs): 18500 KG", "Gross Wt (kgs): N/A"),
        encoding="utf-8")
    email("email_demo_missing_weight", "Please check SI vs draft BL - DEMO-BK-00043",
          compare_body, ["demo_missing_SI.txt", "demo_missing_BL.txt"])

    # 4. An explicit comparison request with nothing attached.
    email("email_demo_no_attachments", "Please check SI vs draft BL - DEMO-BK-00044",
          "Dear Team,\n\nPlease compare the SI and draft BL for DEMO-BK-00044 and "
          "confirm (the attachments appear to have been dropped).\n\nRegards,\nDemo Desk\n",
          [])

    # 5. The BL slot holds a packing list.
    (attachments / "demo_wrongdoc_SI.txt").write_text(SI_HEADING + BASE_SI, encoding="utf-8")
    (attachments / "demo_wrongdoc_BL.txt").write_text(PACKING_LIST, encoding="utf-8")
    email("email_demo_wrong_doc", "Please check SI vs draft BL - DEMO-BK-00045",
          compare_body, ["demo_wrongdoc_SI.txt", "demo_wrongdoc_BL.txt"])

    # 6. A non-BL category: settles OK with no fields and no review.
    email("email_demo_invoice", "Invoice query DEMO-INV-5501",
          "Dear Team,\n\nWe have a query on the invoice amount for DEMO-INV-5501. "
          "Please send the charge breakdown before we release payment.\n\nRegards,\nDemo Desk\n",
          [])

    # 7. Two plausible drafts for the BL role.
    (attachments / "demo_choice_SI.txt").write_text(SI_HEADING + BASE_SI, encoding="utf-8")
    (attachments / "demo_choice_BL_rev1.txt").write_text(BL_HEADING + BASE_BL, encoding="utf-8")
    (attachments / "demo_choice_BL_rev2.txt").write_text(
        BL_HEADING + BASE_BL.replace("Container Count: 4 x 40'HC",
                                     "Container Count: 5 x 40'HC"),
        encoding="utf-8")
    email("email_demo_document_choice", "Please check SI vs draft BL - DEMO-BK-00046",
          compare_body,
          ["demo_choice_SI.txt", "demo_choice_BL_rev1.txt", "demo_choice_BL_rev2.txt"])

    # 8. "Same as consignee" on both sides: a derived notify party.
    (attachments / "demo_reference_SI.txt").write_text(
        SI_HEADING + BASE_SI.replace("Notify Party: MERIDIAN TRADING GMBH",
                                     "Notify Party: SAME AS CONSIGNEE"),
        encoding="utf-8")
    (attachments / "demo_reference_BL.txt").write_text(
        BL_HEADING + BASE_BL.replace("NOTIFY PARTY: MERIDIAN TRADING GMBH",
                                     "NOTIFY PARTY: SAME AS CONSIGNEE"),
        encoding="utf-8")
    email("email_demo_reference", "Please check SI vs draft BL - DEMO-BK-00047",
          compare_body, ["demo_reference_SI.txt", "demo_reference_BL.txt"])

    # 9. PDF pair, with detail rows under a declared total.
    write_text_pdf(attachments / "demo_pdf_SI.pdf", [
        "SHIPPING INSTRUCTION",
        "Shipper: NORTHWIND PAPER EXPORTS PTE LTD",
        "Consignee: MERIDIAN TRADING GMBH",
        "Notify Party: MERIDIAN TRADING GMBH",
        "Port of Loading (POL): SINGAPORE (SGSIN)",
        "Port of Discharge (POD): HAMBURG, GERMANY (DEHAM)",
        "Total Containers: 3 x 40'HC",
        "Gross Weight (KG): 21600 KG",
    ])
    write_text_pdf(attachments / "demo_pdf_BL.pdf", [
        "BILL OF LADING (DRAFT)",
        "SHIPPER: NORTHWIND PAPER EXPORTS PTE LTD",
        "Consignee (Non-Negotiable): MERIDIAN TRADING GMBH",
        "NOTIFY PARTY: MERIDIAN TRADING GMBH",
        "Load Port: SINGAPORE (SGSIN)",
        "Discharge Port: HAMBURG, GERMANY (DEHAM)",
        "",
        "CONTAINER NO.        DESCRIPTION           WEIGHT (KG)",
        "NWPU1000001          40'HC PAPERBOARD      7200",
        "NWPU1000002          40'HC PAPERBOARD      7200",
        "NWPU1000003          40'HC PAPERBOARD      7200",
        "",
        "Container Count: 3 x 40'HC",
        "TOTAL Gross Weight (KG): 21600 KG",
    ])
    email("email_demo_pdf", "Please check SI vs draft BL - DEMO-BK-00048",
          compare_body, ["demo_pdf_SI.pdf", "demo_pdf_BL.pdf"])

    # 10. A scanned BL with no text layer.
    write_text_pdf(attachments / "demo_scan_SI.pdf", [
        "SHIPPING INSTRUCTION",
        "Shipper: NORTHWIND PAPER EXPORTS PTE LTD",
        "Consignee: MERIDIAN TRADING GMBH",
        "Notify Party: MERIDIAN TRADING GMBH",
        "Port of Loading (POL): SINGAPORE (SGSIN)",
        "Port of Discharge (POD): HAMBURG, GERMANY (DEHAM)",
        "Total Containers: 4 x 40'HC",
        "Gross Weight (KG): 18500 KG",
    ])
    write_image_only_pdf(attachments / "demo_scan_BL.pdf")
    email("email_demo_scanned", "Please check SI vs draft BL - DEMO-BK-00049",
          compare_body, ["demo_scan_SI.pdf", "demo_scan_BL.pdf"])

    # 11. A file that announces itself as a PDF and then stops.
    write_text_pdf(attachments / "demo_corrupt_SI.pdf", [
        "SHIPPING INSTRUCTION",
        "Shipper: NORTHWIND PAPER EXPORTS PTE LTD",
        "Consignee: MERIDIAN TRADING GMBH",
        "Notify Party: MERIDIAN TRADING GMBH",
        "Port of Loading (POL): SINGAPORE (SGSIN)",
        "Port of Discharge (POD): HAMBURG, GERMANY (DEHAM)",
        "Total Containers: 4 x 40'HC",
        "Gross Weight (KG): 18500 KG",
    ])
    write_corrupt_pdf(attachments / "demo_corrupt_BL.pdf")
    email("email_demo_corrupt", "Please check SI vs draft BL - DEMO-BK-00050",
          compare_body, ["demo_corrupt_SI.pdf", "demo_corrupt_BL.pdf"])

    # 12. Office formats, with bilingual labels and a multiline address cell.
    write_xlsx(attachments / "demo_office_SI.xlsx", "BL INSTRUCTION", [
        ("Shipper/Exporter", "NORTHWIND PAPER EXPORTS PTE LTD\n12 KEPPEL ROAD, #08-02; SINGAPORE 089057"),
        ("Consignee (Non-Negotiable)", "MERIDIAN TRADING GMBH\nHAFENSTRASSE 44; 20457 HAMBURG, GERMANY"),
        ("Notify", "MERIDIAN TRADING GMBH"),
        ("POL", "SINGAPORE (SGSIN)"),
        ("Port of Discharge", "HAMBURG, GERMANY (DEHAM)"),
        ("Total Containers", "2 x 20'GP"),
        ("Gross Weight (KGS)", 9400),
        ("BOOKING NO.", "DEMO-BK-00051"),
    ])
    write_docx(attachments / "demo_office_BL.docx", "BILL OF LADING (DRAFT)", [
        ("SHIPPER (发货人)", "NORTHWIND PAPER EXPORTS PTE LTD\n12 KEPPEL ROAD, #08-02; SINGAPORE 089057"),
        ("To the Order of (收货人)", "MERIDIAN TRADING GMBH\nHAFENSTRASSE 44; 20457 HAMBURG, GERMANY"),
        ("Notify Party (通知人)", "MERIDIAN TRADING GMBH"),
        ("Port of Loading (装货港)", "SINGAPORE (SGSIN)"),
        ("PORT OF DISCHARGE (卸货港)", "HAMBURG, GERMANY (DEHAM)"),
        ("Container Count (箱数)", "2 x 20'GP"),
        ("Gross Weight毛重(KGS)", "9400"),
        ("Bill of Lading No.", "DEMO-BL-77125"),
    ])
    email("email_demo_office", "Please check SI vs draft BL - DEMO-BK-00051",
          compare_body, ["demo_office_SI.xlsx", "demo_office_BL.docx"])

    # 13. Two competing gross weights the rules refuse to choose between. The
    # model can pick one, citing the document; every gate still applies.
    (attachments / "demo_ambiguous_SI.txt").write_text(SI_HEADING + BASE_SI, encoding="utf-8")
    (attachments / "demo_ambiguous_BL.txt").write_text(
        BL_HEADING + BASE_BL.replace(
            "Gross Wt (kgs): 18500 KG",
            "Gross Wt (kgs): 18500 KG\nGross Weight (KG): 18950 KG",
        ),
        encoding="utf-8")
    email("email_demo_ambiguous_weight", "Please check SI vs draft BL - DEMO-BK-00053",
          compare_body, ["demo_ambiguous_SI.txt", "demo_ambiguous_BL.txt"])

    # 14. An email whose intent no rule recognizes: the configured model's case.
    email("email_demo_needs_model", "DEMO-BK-00052 follow up",
          "Dear Team,\n\nRegarding DEMO-BK-00052 - the paperwork you sent through "
          "yesterday against what the customer originally asked for. Something is "
          "off between the two and we would like your read before we go back to "
          "them.\n\nRegards,\nDemo Desk\n",
          [])

    for record in emails:
        (inbox / f"{record['email_id']}.json").write_text(
            json.dumps(record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )

    (root / "README.md").write_text(
        "# Synthetic demo fixtures\n\n"
        "Independently authored sources for demos and tests. They contain no\n"
        "participant data, which is why documents from this registry are the only\n"
        "ones a public DEMO caller may stream.\n\n"
        "Regenerate with `python scripts/build_demo_fixtures.py`. Expected values\n"
        "live in `tests/`, read off these documents by hand.\n",
        encoding="utf-8",
    )
    print(f"wrote {len(emails)} demo emails and "
          f"{len(list(attachments.iterdir()))} attachments to {root}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    build(parser.parse_args().root)


if __name__ == "__main__":
    main()
